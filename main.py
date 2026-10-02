from tkinter import *
from  tkinter import ttk
from jarvis import ask_jarvis
from student import open_student
from database import create_database
from database import get_dashboard_data
from fee import open_fee
from reports import open_reports
from PIL import Image, ImageTk
from header import add_header
from database import (
    get_dashboard_data,
    get_daily_collection,
    get_month_collection
)
from database import (
    fetch_all_students,
    fetch_all_payments,
    get_daily_history,
    get_month_summary,
    get_dashboard_data
)
from tkinter import filedialog, messagebox
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl import load_workbook
from tkinter import filedialog, messagebox
from database import add_student_import
from database import clear_all_student_data
import os
import sys

create_database()

# ===== Welcome Message =====

root = Tk()
root.title("College Fee Management System")
root.geometry("1600x900")
root.configure(bg="white")

# COLLEGE LOGO
img = Image.open("college_logo.png")
img = img.resize((70, 70))

logo = ImageTk.PhotoImage(img)
top = Frame(root, bg="darkblue", height=90)
top.pack(fill=X)

logo_label = Label(
    top,
    image=logo,
    bg="darkblue"
)
logo_label.pack(side=LEFT, padx=20, pady=10)

title = Label(
    top,
    text="SAI SAMATH DEGREE & Jr COLLEGE",
    font=("Wide Latin", 18, "bold"),
    bg="darkblue",
    fg="white",
    justify=LEFT
)
title.pack(side=LEFT, padx=10)

def resource_path(filename):
    if getattr(sys, "frozen", False):
        base = os.path.dirname(sys.executable)
    else:
        base = os.path.dirname(os.path.abspath(__file__))

    return os.path.join(base, filename)

def student_master():
    open_student()

def fee_collection():
    open_fee()

def dashboard():

    for widget in content.winfo_children():
        widget.destroy()

    total_students, total_fee, paid_fee, balance_fee = get_dashboard_data()

    # ================= TITLE =================

    Label(
        content,
        text="DASHBOARD",
        font=("Arial", 22, "bold"),
        fg="navy",
        bg="white"
    ).pack(pady=10)

    # ================= SUMMARY FRAME =================

    summary = Frame(content, bg="white")
    summary.pack(fill=X, padx=20, pady=10)

    def card(parent, title, value, color):

        box = Frame(
            parent,
            bg=color,
            width=220,
            height=90,
            relief="raised",
            bd=2
        )

        box.pack(side=LEFT, expand=True, padx=10)
        box.pack_propagate(False)

        Label(
            box,
            text=title,
            bg=color,
            fg="white",
            font=("Arial",11,"bold")
        ).pack(pady=(12,2))

        Label(
            box,
            text=value,
            bg=color,
            fg="white",
            font=("Arial",18,"bold")
        ).pack()

    card(summary, "👨‍🎓 Total Students", total_students, "#3498db")
    card(summary, "💰 Total Fee", f"₹{total_fee}", "#9b59b6")
    card(summary, "✅ Paid Fee", f"₹{paid_fee}", "#27ae60")
    card(summary, "⏳ Balance Fee", f"₹{balance_fee}", "#e67e22")

  
    # ================= MAIN AREA ================

    main_frame = Frame(
        content,
        bg="white"
    )

    main_frame.pack(
        fill=BOTH,
        expand=True,
        padx=20,
        pady=10
    )

    # LEFT SIDE
    left_frame = Frame(main_frame, bg="white")
    left_frame.pack(side=LEFT, fill=BOTH, expand=True, padx=5)

    # RIGHT SIDE
    right_frame = Frame(main_frame, bg="white")
    right_frame.pack(side=RIGHT, fill=BOTH, expand=True, padx=5)

    # ================= LEFT : DAILY COLLECTION =================

    daily_box = LabelFrame(
        left_frame,
        text="📅 DAILY COLLECTION",
        font=("Arial",12,"bold"),
        bg="white",
        padx=10,
        pady=10
    )
    daily_box.pack(fill=BOTH, expand=True)

    top_daily = Frame(daily_box, bg="white")
    top_daily.pack(fill=X, pady=5)

    Label(
        top_daily,
        text="Date",
        bg="white",
        font=("Arial",10,"bold")
    ).pack(side=LEFT)

    daily_date = StringVar()

    Entry(
        top_daily,
        textvariable=daily_date,
        width=15
    ).pack(side=LEFT, padx=5)

    daily_total = StringVar(value="₹ 0")

    table = ttk.Treeview(
        daily_box,
        columns=("adm","name","amount"),
        show="headings",
        height=8
    )

    table.heading("adm", text="Admission No")
    table.heading("name", text="Student Name")
    table.heading("amount", text="Amount")

    table.column("adm", width=100)
    table.column("name", width=180)
    table.column("amount", width=100)

    table.pack(fill=BOTH, expand=True, side=LEFT)

    def daily_collection():

        table.delete(*table.get_children())

        rows = get_daily_collection(daily_date.get())

        total = 0

        for row in rows:

            table.insert("", END, values=row)

            total += row[2]

        daily_total.set(f"₹ {total}")

    Button(
        top_daily,
        text="Search",
        bg="blue",
        fg="white",
        command=daily_collection
    ).pack(side=LEFT, padx=10)

    Label(
        daily_box,
        text="Today's Collection",
        bg="white",
        font=("Arial",11,"bold")
    ).pack()

    Label(
        daily_box,
        textvariable=daily_total,
        bg="white",
        fg="green",
        font=("Arial",15,"bold")
    ).pack(pady=5)
    
    # ================= RIGHT : MONTHLY COLLECTION =================

    month_box = LabelFrame(
        right_frame,
        text="📆 MONTHLY COLLECTION",
        font=("Arial",12,"bold"),
        bg="white",
        padx=10,
        pady=10
    )
    month_box.pack(fill=BOTH, expand=True)

    top_month = Frame(month_box, bg="white")
    top_month.pack(fill=X, pady=5)

    Label(
        top_month,
        text="Month",
        bg="white",
        font=("Arial",10,"bold")
    ).pack(side=LEFT)

    month_var = StringVar()

    Entry(
        top_month,
        textvariable=month_var,
        width=12
    ).pack(side=LEFT, padx=5)

    month_total = StringVar(value="₹ 0")

    month_table = ttk.Treeview(
        month_box,
        columns=("adm","name","amount"),
        show="headings",
        height=15
    )

    month_table.heading("adm", text="Admission No")
    month_table.heading("name", text="Student Name")
    month_table.heading("amount", text="Amount")

    month_table.column("adm", width=110)
    month_table.column("name", width=190)
    month_table.column("amount", width=100)

    month_table.pack(fill=BOTH, expand=True, side=RIGHT)

    def month_collection():

        month_table.delete(*month_table.get_children())

        rows = get_month_collection(month_var.get())

        total = 0

        for row in rows:
            month_table.insert("", END, values=row)
            total += row[2]

        month_total.set(f"₹ {total}")

    Button(
        top_month,
        text="Search",
        bg="green",
        fg="white",
        command=month_collection
    ).pack(side=LEFT, padx=10)

    Label(
        month_box,
        text="Monthly Collection",
        bg="white",
        font=("Arial",11,"bold")
    ).pack()

    Label(
        month_box,
        textvariable=month_total,
        bg="white",
        fg="green",
        font=("Arial",15,"bold")
    ).pack(pady=5)

def reports():
    open_reports()

def export_to_excel():

    rows = fetch_all_students()

    if not rows:
        messagebox.showwarning(
            "No Data",
            "No student records found."
        )
        return

    file = filedialog.asksaveasfilename(
        defaultextension=".xlsx",
        filetypes=[("Excel Files", "*.xlsx")],
        title="Save Excel File"
    )

    if not file:
        return

    wb = Workbook()
    ws = wb.active
    ws.title = "Students"
    fee_sheet = wb.create_sheet(title="Fee Collection")

    fee_sheet.append([
        "Payment Date",
        "Admission No",
        "Student Name",
        "Paid Amount"
    ])

    payments = fetch_all_payments()

    for row in payments:
        fee_sheet.append(row)

    daily_sheet = wb.create_sheet(title="Daily Summary")

    daily_sheet.append([
        "Date",
        "No. of Students",
        "Total Collection"
    ])

    daily_rows = get_daily_history()

    for row in daily_rows:
        daily_sheet.append(row)

    month_sheet = wb.create_sheet(title="Monthly Summary")

    month_sheet.append([
        "Month",
        "Total Collection"
    ])

    month_rows = get_month_summary()

    for row in month_rows:
        month_sheet.append(row)

    dash_sheet = wb.create_sheet(title="Dashboard")

    total_students, total_fee, paid_fee, balance_fee = get_dashboard_data()

    dash_sheet.append(["Particular", "Value"])

    dash_sheet.append(["Total Students", total_students])
    dash_sheet.append(["Total Fee", total_fee])
    dash_sheet.append(["Paid Fee", paid_fee])
    dash_sheet.append(["Balance Fee", balance_fee])

    ws.append([
        "Admission No",
        "Student Name",
        "Father Name",
        "Mobile",
        "Course",
        "Year",
        "Total Fee",
        "Paid Fee",
        "Balance Fee"
    ])

    from openpyxl.utils import get_column_letter

    for col in range(1, 10):   # 9 columns

        column = get_column_letter(col)

        max_length = 0

        for cell in ws[column]:

            if cell.row == 1:
                continue

            if cell.value:
                max_length = max(max_length, len(str(cell.value)))

        ws.column_dimensions[column].width = max_length + 4
        
    for row in rows:
        ws.append(row)

    for column_cells in ws.columns:

        length = max(len(str(cell.value)) if cell.value else 0 for cell in column_cells)

        ws.column_dimensions[column_cells[0].column_letter].width = length + 4

    wb.save(file)

    messagebox.showinfo(
        "Success",
        "Excel File Exported Successfully!"
    )

def import_from_excel():

    file = filedialog.askopenfilename(
        title="Select Excel File",
        filetypes=[("Excel Files", "*.xlsx")]
    )

    if not file:
        return

    wb = load_workbook(file)

    ws = wb["Sheet1"]

    count = 0

    for row in ws.iter_rows(min_row=2, values_only=True):

        if not row[0]:
            continue

        add_student_import(
            row[0],  # Admission No
            row[1],  # Student Name
            row[2],  # Father Name
            row[3],  # Mobile
            row[4],  # Course
            row[5],  # Year
            row[6],  # Total Fee
            row[7],  # Paid Fee
            row[8]   # Balance Fee
        )

        count += 1

    messagebox.showinfo(
        "Success",
        f"{count} Students Imported Successfully!"
    )

def clear_imported_data():

    confirm = messagebox.askyesno(
        "Clear Data",
        "All student and payment data will be deleted.\n\n"
        "Do you want to continue?"
    )

    if not confirm:
        return

    clear_all_student_data()

    messagebox.showinfo(
        "Success",
        "Student and payment data cleared successfully!"
    )


    
# ===== Left Menu =====
menu = Frame(root, bg="#2c3e50", width=220)
menu.pack(side=LEFT, fill=Y)

Button(menu,
       text="Student Master",
       font=("Arial",12),
       width=20,
       height=2,
       command=student_master).pack(pady=10)

Button(menu, text="Fee Collection",
font=("Arial",12),
width=20,
height=2,
command=fee_collection).pack(pady=10)

Button(menu, text="Dashboard",
font=("Arial",12),
width=20,
height=2,
command=dashboard).pack(pady=10)

Button(menu, text="Reports",
font=("Arial",12),
width=20,
height=2,
command=reports).pack(pady=10)

Button(
    menu,
    text="Clear Data",
    font=("Arial",12),
    width=20,
    height=2,
    bg="red",
    fg="white",
    command=clear_imported_data
).pack(pady=10)

Button(
    menu,
    text="Export Excel",
    font=("Arial",12),
    width=20,
    height=2,
    command=export_to_excel
).pack(pady=10)

Button(
    menu,
    text="Import Excel",
    font=("Arial",12),
    width=20,
    height=2,
    command=import_from_excel
).pack(pady=10)

Button(menu, text="Exit",
font=("Arial",12),
width=20,
height=2,
command=root.destroy).pack(pady=10)

# ===== Main Area =====

content = Frame(
    root,
    bg="white"
)

content.pack(
    side=LEFT,
    fill=BOTH,
    expand=True
)

# ===== Welcome Message =====

Label(
    content,
    text="Welcome to College Fee Management System",
    font=("Arial",20,"bold"),
    bg="white",
    fg="black"
).pack(pady=(35, 15))


# ================= JARVIS =================

jarvis_main_box = LabelFrame(
    content,
    text="🤖 JARVIS AI",
    font=("Arial",14,"bold"),
    bg="white",
    fg="darkblue",
    padx=25,
    pady=15
)

jarvis_main_box.pack(
    fill=X,
    padx=80,
    pady=10
)


Label(
    jarvis_main_box,
    text="Hello, how can I help you Boss?",
    font=("Arial",15,"bold"),
    fg="darkblue",
    bg="white"
).pack(pady=8)

# ================= JARVIS RESPONSE AREA =================

response_frame = Frame(
    jarvis_main_box,
    bg="white"
)

response_frame.pack(
    fill=BOTH,
    expand=True,
    pady=5
)

jarvis_main_response = Text(
    response_frame,
    height=8,
    width=70,
    font=("Arial", 12),
    bg="white",
    fg="green",
    wrap=WORD,
    relief="solid",
    bd=1
)

jarvis_main_response.pack(
    side=LEFT,
    fill=BOTH,
    expand=True
)

jarvis_scroll = Scrollbar(
    response_frame,
    orient=VERTICAL,
    command=jarvis_main_response.yview
)

jarvis_scroll.pack(
    side=RIGHT,
    fill=Y
)

jarvis_main_response.config(
    yscrollcommand=jarvis_scroll.set
)

# Initial message
jarvis_main_response.insert(
    END,
    "Jarvis: Ready for your command..."
)

jarvis_main_response.config(
    state=DISABLED
)

jarvis_main_input = Entry(
    jarvis_main_box,
    font=("Arial",12),
    width=50
)

jarvis_main_input.pack(
    pady=8
)


def run_main_jarvis():

    command = jarvis_main_input.get().strip()

    if command == "":
        return

    response = ask_jarvis(command)

    if response is None:
        response = "Sorry Boss, I couldn't process your command."

    jarvis_main_response.config(state=NORMAL)

    jarvis_main_response.delete("1.0", END)

    jarvis_main_response.insert(
        END,
        "Jarvis: " + response
    )

    jarvis_main_response.config(state=DISABLED)

    jarvis_main_response.see(END)

    jarvis_main_input.delete(0, END)


Button(
    jarvis_main_box,
    text="ASK JARVIS 🤖",
    font=("Arial",11,"bold"),
    bg="darkblue",
    fg="white",
    padx=25,
    pady=7,
    command=run_main_jarvis
).pack(pady=8)


jarvis_main_input.bind(
    "<Return>",
    lambda event: run_main_jarvis()
)


root.mainloop()

def daily_collection():

    win = Toplevel(root)
    win.title("Daily Collection")
    win.geometry("700x500")
    win.configure(bg="white")
